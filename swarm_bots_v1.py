import json
from docx import Document
from datetime import datetime

print("=== SWARM BOTS v1.0 ===")
print("Generando reporte ISO desde datos de sensores...")
print("")

# Cargar datos del simulador
with open('datos_sensores.json', 'r', encoding='utf-8') as f:
    lecturas = json.load(f)

# Crear documento Word
doc = Document()
doc.add_heading('REPORTE AUTOMÁTICO ISO 9001', 0)
doc.add_paragraph(f'Fecha de generación: {datetime.now().strftime("%d/%m/%Y %H:%M")}')
doc.add_paragraph(f'Total de lecturas analizadas: {len(lecturas)}')

# Sección de resumen
doc.add_heading('Resumen del Enjambre', level=1)
normales = sum(1 for l in lecturas if l['estado'] == "NORMAL")
alertas = sum(1 for l in lecturas if l['estado'] == "ALERTA")
doc.add_paragraph(f'✅ Lecturas normales: {normales}')
doc.add_paragraph(f'⚠️ Alertas detectadas: {alertas}')

# Sección de detalles
doc.add_heading('Detalle de Lecturas', level=1)

for i, lectura in enumerate(lecturas, 1):
    doc.add_paragraph(f'{i}. Robot {lectura["robot_id"]} - {lectura["ubicacion"]}')
    doc.add_paragraph(f'   Sensor: {lectura["sensor"]} = {lectura["valor"]} {lectura["unidad"]}')
    if lectura["estado"] == "ALERTA":
        doc.add_paragraph(f'   ⚠️ NO CONFORMIDAD: Excede límite de {lectura["limite"]} {lectura["unidad"]}', style='Intense Quote')
    doc.add_paragraph('')

# Recomendaciones
doc.add_heading('Recomendaciones Automáticas', level=1)
if alertas > 0:
    doc.add_paragraph('1. Investigar alertas en las próximas 24 horas')
    doc.add_paragraph('2. Programar mantenimiento correctivo')
    doc.add_paragraph('3. Documentar en registro ISO 9001')
else:
    doc.add_paragraph('✅ Sin acciones correctivas requeridas')

# Guardar
nombre_archivo = 'reporte_swarm_bots.docx'
doc.save(nombre_archivo)

print(f'✅ Reporte generado: {nombre_archivo}')
print(f'📊 Lecturas procesadas: {len(lecturas)}')
print(f'⚠️ Alertas: {alertas}')
print("")
print("=== ¡Sistema Swarm Bots funcionando! ===")